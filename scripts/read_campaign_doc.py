import docx
import os
import sys

# Ensure UTF-8 output in Windows PowerShell / CMD
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')

def read_document(filepath="Pokemon_Pen_and_Paper_Kampagne_Akt1.docx"):
    if not os.path.exists(filepath):
        print(f"File '{filepath}' not found.")
        return

    doc = docx.Document(filepath)
    print(f"=== INHALT VON {filepath} ===\n")
    for para in doc.paragraphs:
        if para.text.strip():
            print(para.text)
    
    print("\n--- TABELLEN ---\n")
    for t_idx, table in enumerate(doc.tables):
        print(f"Tabelle #{t_idx + 1}:")
        for row in table.rows:
            row_vals = [cell.text.replace('\n', ' ') for cell in row.cells]
            print(" | ".join(row_vals))
        print("-" * 40)

if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else "Pokemon_Pen_and_Paper_Kampagne_Akt1.docx"
    read_document(path)
